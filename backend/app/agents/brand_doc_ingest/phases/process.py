"""Phase A — process pending brand_documents rows by doc_kind.

For each `brand_documents` row with upload_status='uploaded':
  pdf       -> fitz (PyMuPDF) per-page text → raw_text
  docx      -> python-docx para + table text → raw_text
  text      -> UTF-8 read → raw_text  (Markdown maps here too)
  image     -> OpenAI vision describe (if MIME supported) → ai_description
               sentence-transformers CLIP encode → visual_embedding (512d)
  free_text -> raw_text already set by uploader; no-op extraction

Flips upload_status from 'uploaded' to 'processing'. Leaves the row at
'processing' so the next phase (summarize_embed) can finish it.

On any failure: upload_status='failed' + metadata.error.
"""
from __future__ import annotations

import asyncio
import logging
from pathlib import Path

from ....config import config
from ....db import get_conn
from ....services.clip_service import enabled as clip_enabled, encode_image_paths
from ....services.openai_client import is_vision_supported, vision_describe
from ..._common import materialize_doc
from ...website_ingest.types import CancellationError, PhaseContext, PhaseResult

log = logging.getLogger("archivist.phase.brand_doc.process")


async def _select_pending() -> list[dict]:
    async with get_conn() as conn:
        async with conn.cursor() as cur:
            await cur.execute(
                """
                SELECT id, doc_kind, source_filename, storage_bucket, storage_path,
                       raw_text, ai_description, metadata, mime_type
                  FROM brand_documents
                 WHERE upload_status='uploaded'
                   AND doc_kind <> 'video'   -- videos owned by brand_video_poller
                 ORDER BY id ASC
                """
            )
            return await cur.fetchall()


async def _mark_processing(doc_id: int) -> None:
    async with get_conn() as conn:
        async with conn.cursor() as cur:
            await cur.execute(
                "UPDATE brand_documents SET upload_status='processing' WHERE id=%s",
                (doc_id,),
            )


async def _mark_failed(doc_id: int, reason: str) -> None:
    async with get_conn() as conn:
        async with conn.cursor() as cur:
            await cur.execute(
                "UPDATE brand_documents "
                "   SET upload_status='failed', "
                "       metadata=COALESCE(metadata, '{}'::jsonb) || jsonb_build_object('error', %s::text) "
                " WHERE id=%s",
                (reason[:500], doc_id),
            )


# ---------------------------------------------------------------------------
# Extractors (sync, called via asyncio.to_thread)
# ---------------------------------------------------------------------------

def _extract_pdf_sync(path: Path) -> str:
    import fitz  # pymupdf
    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n\n".join(p for p in parts if p).strip()


def _extract_docx_sync(path: Path) -> str:
    import docx
    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_text_sync(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


# ---------------------------------------------------------------------------
# Per-doc processing
# ---------------------------------------------------------------------------

async def _process_doc(row: dict) -> None:
    doc_id = row["id"]
    kind = row["doc_kind"]
    await _mark_processing(doc_id)

    if kind == "free_text":
        # raw_text already populated when the upload endpoint inserted the row.
        if not (row.get("raw_text") or "").strip():
            raise RuntimeError("free_text doc has empty raw_text")
        log.info("[doc %d] free_text — extraction no-op (%d chars)", doc_id, len(row["raw_text"]))
        return

    if not row.get("storage_path"):
        raise RuntimeError(f"{kind} doc has no storage_path")
    # Materialise to a local path — downloads Supabase-stored docs to a temp file
    # (cleaned up in finally); resolves the legacy local path otherwise.
    path, cleanup = await materialize_doc(row)
    try:
        if not path.exists():
            raise FileNotFoundError(f"storage_path not on disk: {path}")

        if kind == "pdf":
            text = await asyncio.to_thread(_extract_pdf_sync, path)
            if not text:
                raise RuntimeError("pdf extracted to empty text")
            async with get_conn() as conn:
                async with conn.cursor() as cur:
                    await cur.execute(
                        "UPDATE brand_documents SET raw_text=%s WHERE id=%s",
                        (text, doc_id),
                    )
            log.info("[doc %d] pdf → %d chars", doc_id, len(text))
            return

        if kind == "docx":
            text = await asyncio.to_thread(_extract_docx_sync, path)
            if not text:
                raise RuntimeError("docx extracted to empty text")
            async with get_conn() as conn:
                async with conn.cursor() as cur:
                    await cur.execute(
                        "UPDATE brand_documents SET raw_text=%s WHERE id=%s",
                        (text, doc_id),
                    )
            log.info("[doc %d] docx → %d chars", doc_id, len(text))
            return

        if kind == "text":
            text = await asyncio.to_thread(_extract_text_sync, path)
            if not text:
                raise RuntimeError("text file is empty")
            async with get_conn() as conn:
                async with conn.cursor() as cur:
                    await cur.execute(
                        "UPDATE brand_documents SET raw_text=%s WHERE id=%s",
                        (text, doc_id),
                    )
            log.info("[doc %d] text → %d chars", doc_id, len(text))
            return

        if kind == "image":
            # 1. Vision describe (if supported) → ai_description
            desc: str | None = None
            if is_vision_supported(path):
                try:
                    desc = await vision_describe(
                        path,
                        prompt=(
                            "Describe this image in 2-3 sentences focusing on: "
                            "subject, composition, mood, color palette, and what "
                            "it communicates about the brand. Reference any "
                            "visible text."
                        ),
                        max_tokens=200, detail="low",
                    )
                except Exception:
                    log.exception("[doc %d] vision describe failed", doc_id)
            else:
                log.info("[doc %d] image MIME unsupported; skipping vision", doc_id)
                desc = None

            # 2. CLIP encode → visual_embedding
            vec: list[float] | None = None
            if clip_enabled():
                try:
                    vecs = await asyncio.to_thread(encode_image_paths, [path], batch_size=1)
                    vec = vecs[0] if vecs else None
                except Exception:
                    log.exception("[doc %d] CLIP encode failed", doc_id)

            async with get_conn() as conn:
                async with conn.cursor() as cur:
                    await cur.execute(
                        "UPDATE brand_documents "
                        "   SET ai_description=%s, visual_embedding=%s "
                        " WHERE id=%s",
                        (desc, vec, doc_id),
                    )

            if not desc and vec is None:
                raise RuntimeError("image extraction produced no description and no CLIP vec")
            log.info("[doc %d] image → desc=%s, clip=%s", doc_id,
                     len(desc or "") if desc else 0, "yes" if vec else "no")
            return

        raise RuntimeError(f"unknown doc_kind: {kind!r}")
    finally:
        cleanup()


# ---------------------------------------------------------------------------
# Phase entry
# ---------------------------------------------------------------------------

async def run(ctx: PhaseContext) -> PhaseResult:
    rows = await _select_pending()
    if not rows:
        log.info("[run=%s] doc-process: nothing pending", ctx.run_id)
        return PhaseResult(name="doc_process", total=0)

    log.info("[run=%s] doc-process: %d docs to process", ctx.run_id, len(rows))
    processed = 0
    failed = 0

    for i, row in enumerate(rows, 1):
        if ctx.is_cancelled and ctx.is_cancelled():
            raise CancellationError("doc-process cancelled mid-batch")
        try:
            await _process_doc(row)
            processed += 1
        except Exception as e:
            failed += 1
            log.exception("[doc %s] process failed", row["id"])
            await _mark_failed(row["id"], str(e))
        if ctx.report_progress:
            ctx.report_progress("doc_process", processed + failed, len(rows))

    log.info("[run=%s] doc-process done: %d ok, %d failed", ctx.run_id, processed, failed)
    return PhaseResult(
        name="doc_process", processed=processed, failed=failed, total=len(rows),
    )
